#!/usr/bin/env python3
"""
ULTRA CPU Whisper Transcriber
Optimized for systems WITHOUT CUDA/GPU.

Features:
- CPU int8 inference (fastest stable mode)
- Silence filtering (VAD)
- Stable decoding (no hallucination loops)
- Optional forced language
- Clean DOCX output
"""

import sys
from faster_whisper import WhisperModel
from docx import Document

# ================= USER SETTINGS =================
MODEL_SIZE = "small"      # options: tiny, base, small, medium
FORCE_LANGUAGE = None     # "bn", "en", or None for auto
# =================================================


def transcribe(audio_path):
    print("\nLoading model (CPU optimized)...")
    model = WhisperModel(
        MODEL_SIZE,
        device="cpu",
        compute_type="int8"
    )

    print("Starting transcription...\n")

    segments, info = model.transcribe(
        audio_path,
        beam_size=1,
        best_of=1,
        language=FORCE_LANGUAGE,
        task="transcribe",
        temperature=0.0,
        condition_on_previous_text=False,
        vad_filter=True,
        vad_parameters=dict(
            min_silence_duration_ms=500
        ),
        compression_ratio_threshold=2.4,
        log_prob_threshold=-1.0,
        no_speech_threshold=0.6
    )

    print(f"Detected language: {info.language}\n")

    results = []
    for seg in segments:
        start = f"{seg.start:0.1f}s"
        end = f"{seg.end:0.1f}s"
        text = seg.text.strip()
        print(f"{start} - {end}  {text}")
        results.append(text)

    return results


def save_docx(text_list, output_file):
    print("\nSaving DOCX...")
    doc = Document()
    doc.add_heading("Transcription", level=1)

    paragraph = ""
    for line in text_list:
        if len(paragraph) < 500:
            paragraph += " " + line
        else:
            doc.add_paragraph(paragraph.strip())
            paragraph = line

    if paragraph:
        doc.add_paragraph(paragraph.strip())

    doc.save(output_file)
    print(f"Saved: {output_file}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python transcribe_ultra_cpu.py <audio_file>")
        sys.exit(1)

    audio_file = sys.argv[1]
    text = transcribe(audio_file)

    output_file = audio_file.rsplit(".", 1)[0] + ".docx"
    save_docx(text, output_file)


if __name__ == "__main__":
    main()
